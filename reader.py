"""
reader.py — WorkBot File Reader & Duplicate Detector

Reads all CSV (.csv) and Excel (.xlsx) files from the bot's working directory,
detects duplicate rows (same Date + Actual Start + Actual End + Status),
and can optionally write deduplicated versions of those files.

Usage (standalone):
    python reader.py            # report duplicates in all found files
    python reader.py --clean    # report AND write cleaned files (*_clean.csv / *_clean.xlsx)
"""

from __future__ import annotations

import csv
import io
import sys
from pathlib import Path
from typing import NamedTuple

# ── Config ─────────────────────────────────────────────────────
FOLDER = Path(__file__).parent          # same directory as this script
CSV_SUFFIX  = ".csv"
XLSX_SUFFIX = ".xlsx"

# Columns that together define a "unique" data row.
# Adjust if your export columns differ.
DEDUP_KEYS = ("Date", "Actual Start", "Actual End", "Status")


# ── Data structures ────────────────────────────────────────────

class FileReport(NamedTuple):
    path:            Path
    file_type:       str          # "csv" or "xlsx"
    total_rows:      int
    duplicate_rows:  int
    rows:            list[dict]   # all rows (headers preserved as keys)
    unique_rows:     list[dict]   # deduplicated rows (first occurrence kept)


# ── Readers ────────────────────────────────────────────────────

def _read_csv(path: Path) -> list[dict]:
    """Return list of row-dicts from a CSV file (handles UTF-8 BOM)."""
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [dict(row) for row in reader]


def _read_xlsx(path: Path) -> list[dict]:
    """Return list of row-dicts from the first sheet of an XLSX file."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("openpyxl is required: pip install openpyxl")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(h) if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        # Skip completely empty rows
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        result.append({headers[i]: (str(row[i]) if row[i] is not None else "") for i in range(len(headers))})
    return result


# ── Deduplication ───────────────────────────────────────────────

def _dedup_key(row: dict) -> tuple:
    """Build a hashable key from DEDUP_KEYS values."""
    return tuple(str(row.get(k, "")).strip() for k in DEDUP_KEYS)


def _deduplicate(rows: list[dict]) -> tuple[list[dict], list[int]]:
    """
    Return (unique_rows, duplicate_indices).
    First occurrence of each key is kept; subsequent occurrences are duplicates.
    """
    seen: set[tuple] = set()
    unique: list[dict] = []
    dup_indices: list[int] = []

    for i, row in enumerate(rows):
        key = _dedup_key(row)
        if key in seen:
            dup_indices.append(i)
        else:
            seen.add(key)
            unique.append(row)

    return unique, dup_indices


# ── Main scanner ────────────────────────────────────────────────

def scan_folder(folder: Path = FOLDER) -> list[FileReport]:
    """Scan folder for CSV / XLSX files and return a FileReport for each."""
    reports: list[FileReport] = []

    for path in sorted(folder.iterdir()):
        if path.name.startswith("~"):          # skip Excel temp files
            continue
        # Skip previously cleaned files
        if "_clean" in path.stem:
            continue

        if path.suffix.lower() == CSV_SUFFIX:
            rows = _read_csv(path)
            file_type = "csv"
        elif path.suffix.lower() == XLSX_SUFFIX:
            rows = _read_xlsx(path)
            file_type = "xlsx"
        else:
            continue

        unique, dup_indices = _deduplicate(rows)
        reports.append(FileReport(
            path=path,
            file_type=file_type,
            total_rows=len(rows),
            duplicate_rows=len(dup_indices),
            rows=rows,
            unique_rows=unique,
        ))

    return reports


# ── Writers (for --clean mode) ──────────────────────────────────

def _write_csv_clean(report: FileReport) -> Path:
    out = report.path.with_stem(report.path.stem + "_clean")
    if not report.unique_rows:
        return out
    fieldnames = list(report.unique_rows[0].keys())
    with open(out, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(report.unique_rows)
    return out


def _write_xlsx_clean(report: FileReport) -> Path:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise RuntimeError("openpyxl is required: pip install openpyxl")

    out = report.path.with_stem(report.path.stem + "_clean")
    if not report.unique_rows:
        return out

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work History (Cleaned)"

    headers = list(report.unique_rows[0].keys())
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for ri, row in enumerate(report.unique_rows, 2):
        for ci, h in enumerate(headers, 1):
            ws.cell(row=ri, column=ci, value=row.get(h, ""))

    ws.freeze_panes = "A2"
    wb.save(out)
    return out


# ── Reporting ───────────────────────────────────────────────────

def print_report(reports: list[FileReport], clean: bool = False) -> None:
    if not reports:
        print("📂 No CSV or XLSX files found in:", FOLDER)
        return

    total_dups = sum(r.duplicate_rows for r in reports)
    print(f"\n{'='*60}")
    print(f"  WorkBot File Reader — {len(reports)} file(s) scanned")
    print(f"{'='*60}\n")

    for rep in reports:
        icon = "📊" if rep.file_type == "xlsx" else "📄"
        print(f"{icon}  {rep.path.name}")
        print(f"     Rows total   : {rep.total_rows}")
        print(f"     Unique rows  : {len(rep.unique_rows)}")
        print(f"     Duplicates   : {rep.duplicate_rows}", end="")

        if rep.duplicate_rows > 0:
            print("  ⚠️  — showing duplicate keys:")
            # Find and print the duplicate keys for transparency
            seen: set[tuple] = set()
            for row in rep.rows:
                key = _dedup_key(row)
                if key in seen:
                    print(f"       → Date={key[0]}  Start={key[1]}  End={key[2]}  Status={key[3]}")
                else:
                    seen.add(key)
        else:
            print("  ✅")

        if clean and rep.duplicate_rows > 0:
            if rep.file_type == "csv":
                out = _write_csv_clean(rep)
            else:
                out = _write_xlsx_clean(rep)
            print(f"     Cleaned file : {out.name}")

        print()

    if total_dups == 0:
        print("✅ All files are clean — no duplicates found.\n")
    else:
        print(f"⚠️  Total duplicate rows across all files: {total_dups}")
        if not clean:
            print("   Tip: run with --clean to write deduplicated files.\n")
        else:
            print("   Cleaned files have been written.\n")


# ── Telegram helper ─────────────────────────────────────────────

def build_telegram_report(reports: list[FileReport]) -> str:
    """
    Return a Markdown-formatted string suitable for a Telegram message.
    Called from bot.py to let the bot report on local files.
    """
    if not reports:
        return "📂 No CSV or XLSX files found in the bot folder."

    lines = ["*📂 File Scan Report*\n"]
    for rep in reports:
        icon = "📊" if rep.file_type == "xlsx" else "📄"
        dups = rep.duplicate_rows
        status = f"⚠️ *{dups} duplicate(s)*" if dups else "✅ Clean"
        lines.append(
            f"{icon} `{rep.path.name}`\n"
            f"   Rows: {rep.total_rows} | Unique: {len(rep.unique_rows)} | {status}"
        )
        if dups:
            seen: set[tuple] = set()
            shown = 0
            for row in rep.rows:
                key = _dedup_key(row)
                if key in seen and shown < 5:
                    lines.append(f"   ↳ `{key[0]}` start=`{key[1]}` status=`{key[3]}`")
                    shown += 1
                else:
                    seen.add(key)
            if dups > 5:
                lines.append(f"   ↳ _…and {dups - 5} more_")
        lines.append("")

    total_dups = sum(r.duplicate_rows for r in reports)
    if total_dups:
        lines.append(f"⚠️ Total duplicates: *{total_dups}*")
        lines.append("Use /cleanfiles to write deduplicated copies.")
    else:
        lines.append("✅ All files are duplicate-free.")

    return "\n".join(lines)


# ── CLI entry point ─────────────────────────────────────────────

if __name__ == "__main__":
    clean_mode = "--clean" in sys.argv
    reports = scan_folder()
    print_report(reports, clean=clean_mode)
